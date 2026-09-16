"""
Модуль М1: Загрузка файлов

Действия:
- Чтение xlsx/xlsm/xls/CSV
- Ленивая загрузка стилей (FR-0103)
- Сохранение цветов и форматов
"""

from pathlib import Path
from typing import List, Dict, Any, Optional
import logging

logger = logging.getLogger('ExcelStudioHR.modules.m01_load')


class SheetModel:
    """Модель листа в памяти."""
    
    def __init__(self, name: str):
        self.name = name
        self.data: List[List[Any]] = []  # Матрица значений
        self.styles: List[List[Dict]] = []  # Матрица стилей
        self.column_widths: List[float] = []
        self.row_heights: List[float] = []
        self.merges: List[str] = []  # Диапазоны объединений
    
    @property
    def row_count(self) -> int:
        return len(self.data)
    
    @property
    def col_count(self) -> int:
        if not self.data:
            return 0
        return max(len(row) for row in self.data)


class FileModel:
    """Модель файла Excel."""
    
    def __init__(self, path: Path):
        self.path = path
        self.sheets: Dict[str, SheetModel] = {}
        self.styles_loaded: bool = False  # Флаг полной загрузки стилей
    
    def add_sheet(self, name: str) -> SheetModel:
        sheet = SheetModel(name)
        self.sheets[name] = sheet
        return sheet


def read_any_to_models(file_path: Path) -> Optional[FileModel]:
    """
    Диспетчер чтения файлов по расширению (FR-0101).
    """
    ext = file_path.suffix.lower()
    
    if ext in ('.xlsx', '.xlsm'):
        return _read_openpyxl(file_path, read_styles=False)
    elif ext == '.xls':
        logger.warning(f'Чтение .xls требует xlrd: {file_path}')
        return _read_xls(file_path)
    elif ext == '.csv':
        return _read_csv(file_path)
    else:
        logger.error(f'Неподдерживаемый формат: {ext}')
        return None


def _read_openpyxl(file_path: Path, read_styles: bool = False) -> Optional[FileModel]:
    """Чтение xlsx/xlsm через openpyxl (FR-0102, FR-0103)."""
    try:
        from openpyxl import load_workbook
        
        wb = load_workbook(filename=str(file_path), data_only=True, read_only=not read_styles)
        file_model = FileModel(file_path)
        
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            sheet_model = file_model.add_sheet(sheet_name)
            
            # Читаем значения всех листов всегда (FR-0103a)
            sheet_model.data = []
            for row in ws.iter_rows(values_only=True):
                sheet_model.data.append(list(row))
            
            # Стили загружаются лениво
            if read_styles:
                _load_styles_full(ws, sheet_model)
            else:
                _load_styles_lazy(ws, sheet_model, limit=50)
            
            # Ширины столбцов
            for col_dim in ws.column_dimensions.values():
                idx = _col_letter_to_idx(col_dim.column) - 1
                if idx >= 0:
                    while len(sheet_model.column_widths) <= idx:
                        sheet_model.column_widths.append(8.5)
                    sheet_model.column_widths[idx] = col_dim.width or 8.5
            
            # Высоты строк
            for row_dim in ws.row_dimensions.values():
                idx = row_dim.row - 1
                if idx >= 0:
                    while len(sheet_model.row_heights) <= idx:
                        sheet_model.row_heights.append(15.0)
                    sheet_model.row_heights[idx] = row_dim.height or 15.0
            
            # Объединения
            for merge_range in ws.merged_cells.ranges:
                sheet_model.merges.append(str(merge_range))
        
        wb.close()
        file_model.styles_loaded = read_styles
        
        return file_model
        
    except Exception as e:
        logger.error(f'Ошибка чтения {file_path}: {e}', exc_info=True)
        return None


def _load_styles_lazy(ws, sheet_model: SheetModel, limit: int = 50):
    """Ленивая загрузка стилей (только шапка и первые строки)."""
    sheet_model.styles = []
    for row_idx, row in enumerate(ws.iter_rows()):
        if row_idx >= limit:
            break
        style_row = []
        for cell in row:
            style = _extract_cell_style(cell)
            style_row.append(style)
        sheet_model.styles.append(style_row)


def _load_styles_full(ws, sheet_model: SheetModel):
    """Полная загрузка стилей."""
    sheet_model.styles = []
    for row in ws.iter_rows():
        style_row = []
        for cell in row:
            style = _extract_cell_style(cell)
            style_row.append(style)
        sheet_model.styles.append(style_row)


def _extract_cell_style(cell) -> Dict:
    """Извлечение стиля ячейки."""
    return {
        'font': {
            'name': cell.font.name,
            'size': cell.font.size,
            'bold': cell.font.bold,
            'italic': cell.font.italic,
            'color': _color_to_hex(cell.font.color) if cell.font.color else '#000000',
        },
        'fill': {
            'fgColor': _color_to_hex(cell.fill.fgColor) if cell.fill.fgColor else None,
        },
        'number_format': cell.number_format,
        'alignment': {
            'horizontal': cell.alignment.horizontal,
            'vertical': cell.alignment.vertical,
        }
    }


def _color_to_hex(color) -> Optional[str]:
    """Преобразование цвета в HEX."""
    if color is None:
        return None
    if hasattr(color, 'rgb') and color.rgb:
        rgb = color.rgb
        if isinstance(rgb, str):
            return '#' + rgb[-6:]
    return None


def _col_letter_to_idx(col_letter: str) -> int:
    """Преобразование буквы столбца в индекс (A=1, B=2, ..., AA=27)."""
    result = 0
    for char in col_letter.upper():
        result = result * 26 + (ord(char) - ord('A') + 1)
    return result


def _read_xls(file_path: Path) -> Optional[FileModel]:
    """Чтение старых .xls файлов через xlrd."""
    try:
        import xlrd
        
        wb = xlrd.open_workbook(str(file_path))
        file_model = FileModel(file_path)
        
        for sheet_name in wb.sheet_names():
            ws = wb.sheet_by_name(sheet_name)
            sheet_model = file_model.add_sheet(sheet_name)
            
            sheet_model.data = []
            for row_idx in range(ws.nrows):
                row = ws.row_values(row_idx)
                sheet_model.data.append(row)
        
        return file_model
        
    except ImportError:
        logger.error(f'xlrd не установлен для чтения .xls: {file_path}')
        return None
    except Exception as e:
        logger.error(f'Ошибка чтения .xls {file_path}: {e}', exc_info=True)
        return None


def _read_csv(file_path: Path) -> Optional[FileModel]:
    """Чтение CSV файлов."""
    try:
        import csv
        
        file_model = FileModel(file_path)
        sheet_model = file_model.add_sheet(file_path.stem)
        
        with open(file_path, 'r', encoding='utf-8-sig') as f:
            reader = csv.reader(f)
            sheet_model.data = [list(row) for row in reader]
        
        return file_model
        
    except Exception as e:
        logger.error(f'Ошибка чтения CSV {file_path}: {e}', exc_info=True)
        return None


__all__ = ['SheetModel', 'FileModel', 'read_any_to_models']
